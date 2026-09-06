from pathlib import Path
from openpyxl import load_workbook
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment


BASE_DIR = Path(__file__).resolve().parent.parent
REPORT_FILE = BASE_DIR / "data" / "output" / "Sales_Report.xlsx"


def style_header(ws):
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)

    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")


def create_dashboard(wb):
    if "Dashboard" in wb.sheetnames:
        del wb["Dashboard"]

    ws = wb.create_sheet("Dashboard", 0)

    ws.merge_cells("A1:J2")
    ws["A1"] = "SALES DASHBOARD"
    ws["A1"].font = Font(
        bold=True,
        size=24,
        color="FFFFFF"
    )
    ws["A1"].fill = PatternFill(
        "solid",
        fgColor="1F4E78"
    )
    ws["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    ws.merge_cells("A3:J3")
    ws["A3"] = "Excel Sales Automation — Portfolio Dashboard"
    ws["A3"].alignment = Alignment(horizontal="center")

    raw_ws = wb["Raw_Data"]
    region_ws = wb["By_Region"]
    product_ws = wb["By_Product"]
    month_ws = wb["By_Month"]

    total_sales = sum(
        cell.value or 0
        for cell in raw_ws["H"][1:]
    )

    total_quantity = sum(
        cell.value or 0
        for cell in raw_ws["F"][1:]
    )

    orders = len(
        {
            cell.value
            for cell in raw_ws["A"][1:]
            if cell.value is not None
        }
    )

    cards = [
        ("B5", "TOTAL SALES", total_sales),
        ("E5", "TOTAL ORDERS", orders),
        ("H5", "TOTAL QUANTITY", total_quantity),
    ]

    for cell, label, value in cards:
        ws[cell] = label
        ws[cell].font = Font(
            bold=True,
            color="1F4E78"
        )

        value_cell = ws.cell(
            row=ws[cell].row + 1,
            column=ws[cell].column
        )

        value_cell.value = value
        value_cell.font = Font(
            bold=True,
            size=18
        )

    # Monthly chart
    line_chart = LineChart()
    line_chart.title = "Monthly Sales Trend"
    line_chart.height = 8
    line_chart.width = 12

    data = Reference(
        month_ws,
        min_col=2,
        min_row=1,
        max_row=month_ws.max_row
    )

    categories = Reference(
        month_ws,
        min_col=1,
        min_row=2,
        max_row=month_ws.max_row
    )

    line_chart.add_data(
        data,
        titles_from_data=True
    )

    line_chart.set_categories(categories)

    ws.add_chart(
        line_chart,
        "A10"
    )

    # Region chart
    bar_chart = BarChart()
    bar_chart.type = "bar"
    bar_chart.title = "Sales by Region"
    bar_chart.height = 8
    bar_chart.width = 12

    data = Reference(
        region_ws,
        min_col=2,
        min_row=1,
        max_row=region_ws.max_row
    )

    categories = Reference(
        region_ws,
        min_col=1,
        min_row=2,
        max_row=region_ws.max_row
    )

    bar_chart.add_data(
        data,
        titles_from_data=True
    )

    bar_chart.set_categories(categories)

    ws.add_chart(
        bar_chart,
        "F10"
    )

    # Top products
    ws["A27"] = "TOP PRODUCTS"
    ws["A27"].font = Font(
        bold=True,
        color="FFFFFF"
    )
    ws["A27"].fill = PatternFill(
        "solid",
        fgColor="1F4E78"
    )

    ws["B29"] = "Product"
    ws["C29"] = "Sales"

    for i, row in enumerate(
        product_ws.iter_rows(
            min_row=2,
            max_row=min(
                product_ws.max_row,
                6
            ),
            values_only=True
        ),
        start=30
    ):
        ws[f"B{i}"] = row[0]
        ws[f"C{i}"] = row[2]


def format_sheets(wb):
    for sheet_name in [
        "Raw_Data",
        "By_Region",
        "By_Product",
        "By_Month"
    ]:
        ws = wb[sheet_name]

        style_header(ws)

        ws.freeze_panes = "A2"

        for column in ws.columns:
            max_length = 0

            column_letter = column[0].column_letter

            for cell in column:
                if cell.value is not None:
                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )

            ws.column_dimensions[
                column_letter
            ].width = min(
                max_length + 2,
                25
            )


def generate_dashboard():
    if not REPORT_FILE.exists():
        raise FileNotFoundError(
            f"Report topilmadi: {REPORT_FILE}"
        )

    wb = load_workbook(REPORT_FILE)

    create_dashboard(wb)
    format_sheets(wb)

    wb.save(REPORT_FILE)

    print("Dashboard successfully created:")
    print(REPORT_FILE)


if __name__ == "__main__":
    generate_dashboard()